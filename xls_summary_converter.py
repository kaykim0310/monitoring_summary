# -*- coding: utf-8 -*-
"""측정 프로그램(MES)이 내보낸 엑셀 결과를 분포실태 텍스트로 변환.

지원 형식:
1) MES에서 F11로 받은 원본 엑셀(.xls) — 확장자만 .xls인 HTML 표 파일
2) 위 2개를 시트 2개로 합쳐 저장한 통합 문서(.xlsx)

셀 값이 온전히 저장되어 있어 PDF와 달리 줄바꿈 잘림·컬럼 추측이 필요 없고,
'물질분류' 열이 있으면 그 분류를 그대로 사용한다.

- 소음제외 표: 헤더에 '유해인자' 열 존재 (물질분류 열 포함)
- 소음 표    : 헤더에 '발생형태' 열 존재 (유해인자 열 없음 -> 전부 '소음')
표 종류는 헤더로 자동 판별하므로 파일/시트 순서는 무관하다.
"""
import io
import re
from collections import defaultdict
from html.parser import HTMLParser

from pdf_summary_converter import classify_factor, render_summary

# 보고서의 물질분류 표기 -> 분포실태 표시 카테고리
XLS_CATEGORY_MAP = {
    "분진": "분진류",
    "분진류": "분진류",
    "금속류": "금속류",
    "유기화합물": "유기화합물",
    "산 및 알카리류": "산 및 알칼리류",
    "산 및 알칼리류": "산 및 알칼리류",
    "가스상물질류": "가스상 물질류",
    "가스상 물질류": "가스상 물질류",
    "금속가공유": "금속가공유",
    "물리적인자": "물리적인자",
    "허가대상 유해물질": "허가대상 유해물질",
}


class _TableParser(HTMLParser):
    """HTML 문서의 <table>들을 행렬(문자열)로 추출."""

    def __init__(self):
        super().__init__()
        self.tables = []
        self._rows = None
        self._row = None
        self._cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._rows = []
        elif tag == "tr" and self._rows is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = ""
        elif tag == "br" and self._cell is not None:
            self._cell += "\n"

    def handle_endtag(self, tag):
        if tag == "table" and self._rows is not None:
            if self._rows:
                self.tables.append(self._rows)
            self._rows = None
        elif tag == "tr" and self._row is not None:
            self._rows.append(self._row)
            self._row = None
        elif tag in ("td", "th") and self._cell is not None:
            self._row.append(self._cell.strip())
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell += data


def _decode(data):
    for enc in ("utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, AttributeError):
            continue
    return data.decode("utf-8", errors="replace")


def parse_html_tables(data):
    """bytes(또는 str) HTML에서 표 목록 추출."""
    txt = _decode(data) if isinstance(data, (bytes, bytearray)) else data
    p = _TableParser()
    p.feed(txt)
    return p.tables


def _cell_to_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx_tables(data):
    """진짜 엑셀 통합 문서(.xlsx)의 모든 시트를 표 목록으로 추출."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [[_cell_to_str(c) for c in row]
                for row in ws.iter_rows(values_only=True)]
        if rows:
            tables.append(rows)
    wb.close()
    return tables


def parse_any_tables(data):
    """파일 내용으로 형식을 판별해 표 목록 추출 (xlsx zip 시그니처 = PK)."""
    if isinstance(data, (bytes, bytearray)) and data[:2] == b"PK":
        return parse_xlsx_tables(bytes(data))
    return parse_html_tables(data)


def _find_data_table(tables):
    """'공정명' 헤더가 있는 데이터 표와 헤더 인덱스 맵을 찾는다."""
    for t in tables:
        for hi, row in enumerate(t[:5]):
            if "공정명" in row and "단위작업장소" in row:
                idx = {name: i for i, name in enumerate(row) if name}
                return t[hi + 1:], idx
    return None, None


def _detect_kind(idx):
    if "유해인자" in idx:
        return "chem"
    if "발생형태" in idx:
        return "noise"
    return None


def classify_from_xls(factor_name, xls_category):
    """보고서의 물질분류 값을 우선 사용, 없으면 사전/키워드 분류로 폴백."""
    cat = XLS_CATEGORY_MAP.get(re.sub(r"\s+", "", xls_category or "").replace("알카리", "알칼리").strip())
    if cat:
        return cat
    cat = XLS_CATEGORY_MAP.get((xls_category or "").strip())
    if cat:
        return cat
    return classify_factor(factor_name) or "기타"


def extract_jobs_from_xls(file_datas):
    """HTML-xls 파일들(bytes 목록)에서 jobs 구조를 만든다.

    반환: (company_info, jobs)  — pdf_summary_converter.render_summary와 호환.
    """
    company_info = {}
    jobs = defaultdict(list)
    unit_index = {}   # (공정, 단위) -> entry

    def get_entry(group, unit, workers, form):
        key = (group, unit)
        if key not in unit_index:
            entry = {
                "group": group,
                "job_content": unit,
                "name_parts": [unit] if unit else [],
                "factors": defaultdict(dict),  # 이름 -> None (입력 순서 보존)
                "workers": workers or "",
                "work_form": set(),
            }
            if form:
                entry["work_form"].add(form)
            jobs[group].append(entry)
            unit_index[key] = entry
        else:
            entry = unit_index[key]
            if workers and not entry["workers"]:
                entry["workers"] = workers
            if form:
                entry["work_form"].add(form)
        return entry

    parsed = []
    shell_detected = False
    for data in file_datas:
        # Excel '웹 보관 파일' 껍데기 감지: 데이터가 별도 .files 폴더에 분리된 형태
        if isinstance(data, (bytes, bytearray)) and b"Excel Workbook Frameset" in bytes(data[:2000]):
            shell_detected = True
            continue
        # xlsx는 시트별로, HTML-xls는 표별로 각각 판별
        for table in parse_any_tables(data):
            rows, idx = _find_data_table([table])
            if rows is None:
                continue
            kind = _detect_kind(idx)
            if kind:
                parsed.append((kind, rows, idx))

    kinds = {k for k, _, _ in parsed}
    if "chem" not in kinds:
        msg = "소음제외(유해인자) 표를 찾지 못해 유해인자를 읽을 수 없습니다."
        if shell_detected:
            msg += (" 업로드된 파일 중 데이터가 없는 '껍데기' 파일이 있습니다"
                    " (Excel에서 '웹 페이지'로 재저장되면 용량 10KB 정도의"
                    " 껍데기만 남습니다). MES에서 F11로 받은 원본 파일을"
                    " 수정 없이 그대로 올려주세요.")
        else:
            msg += (" MES에서 F11로 받은 소음제외 원본 파일(.xls)을 함께"
                    " 올렸는지 확인해 주세요.")
        raise ValueError(msg)

    # 소음제외(chem) 먼저 처리해 단위 순서를 잡고, 소음은 뒤에 합류
    parsed.sort(key=lambda x: 0 if x[0] == "chem" else 1)

    for kind, rows, idx in parsed:
        gi = idx.get("공정명")
        ui = idx.get("단위작업장소")
        wi = idx.get("근로자수")
        fmi = idx.get("근로형태")
        fi = idx.get("유해인자")
        ci = idx.get("물질분류")
        co = idx.get("공장명")

        last_group = ""
        last_unit = ""
        for r in rows:
            def cell(i):
                return r[i].strip() if i is not None and i < len(r) else ""

            group = cell(gi)
            if not group:
                continue  # 페이지 나눔 등 빈 행
            unit = cell(ui)
            if not unit:
                # 단위작업장소가 빈 행은 직전 행(같은 공정)의 단위에 귀속
                unit = last_unit if group == last_group else ""
            last_group, last_unit = group, unit

            if co is not None and "name" not in company_info and cell(co):
                company_info["name"] = cell(co)

            entry = get_entry(group, unit, cell(wi), cell(fmi))
            if kind == "chem":
                factor = cell(fi)
                if factor:
                    entry["factors"][classify_from_xls(factor, cell(ci))][factor] = None
            else:
                entry["factors"]["물리적인자"]["소음"] = None

    return company_info, jobs


def convert_xls_to_txt(file_datas):
    """HTML-xls 파일 bytes 목록 -> 분포실태 텍스트."""
    company_info, jobs = extract_jobs_from_xls(file_datas)
    return render_summary(company_info, jobs)
